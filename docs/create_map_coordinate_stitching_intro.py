from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"D:\1\yuanshibanben\PointCloud20260514\SLAM地图坐标系统一与地图拼接介绍.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_table_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def remove_paragraph_borders(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(7)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(text))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        set_run_font(run, size=15 if level == 1 else 12, bold=True)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, title in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "244062")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(title), size=9.5, bold=True, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[i])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        keep_table_row_together(row)
        cells = row.cells
        for i, value in enumerate(values):
            cell = cells[i]
            if row_index % 2:
                set_cell_shading(cell, "F3F6FA")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(str(value)), size=9.3)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.35)
section.right_margin = Cm(2.35)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)
for style_name in ("Title", "Heading 1", "Heading 2"):
    style = styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(10)
remove_paragraph_borders(styles["Title"])
remove_paragraph_borders(title)
set_run_font(title.add_run("SLAM 地图坐标系统一与地图拼接介绍"), size=22, bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
set_run_font(subtitle.add_run("会议介绍稿  大白话版"), size=11, color="555555")

add_body(doc, "这份文档主要回答三个问题：每一份点云地图里的坐标是怎么来的，两份地图为什么能够放到同一个位置，以及程序如何避免一次把整张大地图都塞进内存。当前方案的核心可以概括为一句话：先用经纬度把多份地图放进同一个 ENU 坐标系，再用点云重合情况做小范围修正，最后把结果切成小块，界面只加载眼前需要看的区域。")

add_heading(doc, "一 先说结论", 1)
add_body(doc, "目前新生成的 .slammap 不是一堆没有位置含义的散点。每张地图都保存了地理锚点，并且新版文件还保存了每个关键帧附近的 GNSS 和 INS 参考位姿。程序加载多张地图时，会先把它们换算到第一张地图所使用的统一 ENU 坐标系下。这样，同一个码头、岸线或泊位在不同文件中，理论上应该落在相同的米制坐标附近。")
add_body(doc, "经纬度负责告诉程序“这一张地图大概应该放在哪里”，点云配准负责在证据足够时再做几米以内的小修正。为了防止码头里大量相似的平行岸线把地图吸到错误位置，当前自动精配准默认只修正 X 和 Y 平移，不随意旋转整张地图；质量不够时宁愿保留 ENU 粗对齐结果，也不强行拼。")

add_heading(doc, "二 项目里几种坐标到底是什么", 1)
add_table(doc,
          ["坐标名称", "大白话理解", "主要用途"],
          [
              ["雷达坐标系", "以当前雷达为中心，点表示物体在雷达前后左右多远。船一移动，这套坐标也跟着移动。", "接收每一帧原始点云、实时检测"],
              ["船体坐标系", "以船自身为中心，统一描述船头、船尾、左舷和右舷。雷达安装位置通过外参换到这里。", "船体姿态、不同传感器关系"],
              ["单次 SLAM 地图坐标系", "某一次建图过程使用的固定地图坐标。关键帧位姿负责把当时的雷达点放到这张地图里。", "生成一份连续的局部点云地图"],
              ["WGS84 经纬度", "GNSS 给出的经度、纬度和高程，适合说明地球上的绝对位置，但不适合直接拿“度”做点云距离计算。", "提供地理锚点"],
              ["ENU 坐标系", "把某个经纬度位置定为原点，用米表示向东 E、向北 N、向上 U 的距离。", "多文件统一坐标和地图拼接"],
          ], [3.0, 8.1, 4.9])
add_body(doc, "最容易混淆的是“单次 SLAM 地图坐标”和“统一 ENU”。SLAM 擅长保证短时间内前后两帧接得顺，但会随着时间积累一些漂移；ENU 由 GNSS 和 INS 提供地理方向与位置参考，能让不同航次知道自己在同一片真实区域。两者不是互相替代，而是各管一部分。")

add_heading(doc, "三 一份点云地图是怎么生成的", 1)
add_number(doc, "两个雷达的数据先按照时间和标定关系进行处理，形成当前时刻可用于 SLAM 的点云。")
add_number(doc, "SLAM 根据相邻点云之间的几何关系，估计这一帧雷达相对于地图的位置和方向。")
add_number(doc, "程序不是把每一帧都永久保存，而是按条件挑选关键帧。每个关键帧保存局部点云和当时的 SLAM 位姿。")
add_number(doc, "新版地图还给关键帧保存时间邻近的 GNSS 和 INS 雷达 ENU 参考位姿，作为后续检查和校正漂移的依据。")
add_number(doc, "点击保存地图后，地理锚点、关键帧位姿、地理参考和局部点云一起写入 .slammap 文件。")

add_heading(doc, "四 两份地图具体怎么拼到一起", 1)
add_heading(doc, "第一步 先校正每张地图内部的慢慢跑偏", 2)
add_body(doc, "新版 .slammap 会比较“SLAM 认为关键帧在哪里”和“GNSS、INS 参考认为它在哪里”。程序不会把某一个瞬间的 GNSS 数值生硬地盖到 SLAM 上，而是收集多个有效关键帧，剔除明显跳点，再把修正量平滑地分配到整段轨迹。简单说，就是保留 SLAM 的局部连续性，同时用地理参考把长期漂移拉回来。")

add_heading(doc, "第二步 用地理锚点统一到同一个 ENU", 2)
add_body(doc, "多选地图时，第一张地图作为基准。程序把其他地图的经纬度锚点先换成地球直角坐标，再计算它们相对基准锚点向东、向北、向上的米制距离。然后把其他地图整体平移到第一张地图的 ENU 坐标中。这里使用的是 WGS84 椭球换算，不是简单拿“经纬度差乘一个固定米数”。")

add_heading(doc, "第三步 用重合点云做保守的小修正", 2)
add_body(doc, "GNSS 本身可能有小偏差，因此两张地图放进统一 ENU 后，固定结构还可能差一两米。程序会从岸线、建筑、桥梁、泊位等被多次观察到的稳定区域中寻找重合关系。当前自动模式默认只允许修正 X 和 Y 平移，最大修正范围控制在约 3 米，不自动改变 ENU 航向。")
add_body(doc, "一个候选修正只有在双向匹配率、误差和改善幅度都达到要求时才会被接受。如果点太少、重合不足、结果贴着搜索边界，或者修正反而没有明显变好，程序就放弃这次精配准，回退到经纬度锚点给出的 ENU 位置。")

add_heading(doc, "第四步 显示时做去重但不破坏原文件", 2)
add_body(doc, "两张地图对齐后，重合区域会包含两份点。界面显示阶段使用约 0.3 米的三维体素格做去重：同一个小立方格里已经有点时，后加载的重复点不再显示；高度不同的真实点仍会保留。这个处理只影响显示缓存，不会删除 .slammap 里的原始关键帧和点云。")

flow_heading = add_heading(doc, "五 整个流程用一句流程图表示", 1)
flow = doc.add_table(rows=1, cols=1)
flow.alignment = WD_TABLE_ALIGNMENT.CENTER
flow.style = "Table Grid"
cell = flow.cell(0, 0)
keep_table_row_together(flow.rows[0])
set_cell_margins(cell, 90, 180, 90, 180)
set_cell_shading(cell, "F3F6FA")
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing = 1.0
set_run_font(p.add_run(
    "雷达点云和 GNSS INS\n↓\n"
    "生成关键帧  保存 SLAM 位姿和 ENU 地理参考\n↓\n"
    "单张地图内部漂移校正\n↓\n"
    "按地理锚点放入第一张地图的统一 ENU\n↓\n"
    "重合稳定结构的小范围平移精配准\n↓\n"
    "三维体素显示去重\n↓\n"
    "生成只读分块缓存并按当前视野显示"), size=9.5, bold=True)
doc.add_paragraph().paragraph_format.space_after = Pt(1)

add_heading(doc, "六 为什么还要做只读分块动态显示", 1)
add_body(doc, "如果港区越来越大，点云越来越密，一次把所有点都加载到内存和显存，程序早晚会变慢甚至跑不动。因此当前实现不修改原始 .slammap，而是在第一次加载后生成一套可以随时删除、重新生成的只读分块缓存。")
add_table(doc,
          ["项目", "当前设置", "它解决什么问题"],
          [
              ["地图分块", "50 米乘 50 米", "把大地图拆成许多可独立读取的小块"],
              ["三级精度", "1.0 米  0.3 米  0.1 米", "看全局时少画点，靠近查看时再用高精度"],
              ["预加载", "视野周围提前加载一圈", "拖动地图时减少黑块和等待"],
              ["自动卸载", "离开视野两圈后释放", "控制内存占用"],
              ["内存预算", "512 MB", "不让离线地图无限占用 RAM"],
              ["GPU 点数预算", "150 万点", "避免显卡一次绘制过多点"],
          ], [3.0, 4.3, 8.7])
add_body(doc, "所以，加载地图有两个阶段。第一次选择某一组 .slammap 时，需要完整读取、统一 ENU、拼接并生成 Tile 缓存，等待时间会长一些；第二次加载同一组文件时，如果源文件、顺序和参数都没有变化，就直接命中缓存。进入界面后，鼠标拖到哪里就加载哪里的 Tile，缩放时自动选择合适精度。")

add_heading(doc, "七 目前方案能解决什么", 1)
add_bullet(doc, "让不同文件中的点云使用同一个东、北、天方向和同一个米制基准。")
add_bullet(doc, "利用逐关键帧地理参考，减少单张地图内部随着航行逐渐积累的漂移。")
add_bullet(doc, "用严格门槛限制自动点云配准，避免相似岸线导致整张地图错误旋转。")
add_bullet(doc, "对重合区域做显示去重，减少红绿双影和重复覆盖。")
add_bullet(doc, "通过只读分块、LOD 和缓存预算，支持更大范围地图的浏览。")

add_heading(doc, "八 目前不能靠拼接程序凭空解决什么", 1)
add_bullet(doc, "如果原始 GNSS 或 INS 长时间整体偏移，统一 ENU 也会继承这部分绝对误差。")
add_bullet(doc, "如果两份地图几乎没有共同的固定结构，点云精配准没有足够证据，只能使用地理锚点。")
add_bullet(doc, "旧版 v2.0 地图没有逐关键帧地理参考，无法像新版 v2.1 一样校正整段内部漂移。")
add_bullet(doc, "单次 SLAM 本身已经出现明显分叉或局部形变时，简单整体平移不能把内部每一段都修好。")
add_bullet(doc, "水面、移动船只和重复排列的泊位结构不适合作为唯一配准依据。")

add_heading(doc, "九 建议后续重点优化", 1)
add_table(doc,
          ["优先级", "优化内容", "大白话说明"],
          [
              ["高", "增加拼接质量可视化", "在界面或日志中明确显示锚点偏差、匹配率、误差、修正量和最终是否接受，方便判断是文件问题还是算法问题。"],
              ["高", "建立可靠静态结构掩码", "尽量使用岸线、建筑和固定设施，减少水面、移动船和临时障碍物对配准的影响。"],
              ["高", "统一传感器时间基准", "重点检查雷达、GNSS 和 INS 的时间差；时间错了，位置再准也会把点放错地方。"],
              ["中", "加入人工控制点兜底", "自动结果证据不足时，可让用户在两张地图上各点选几个相同位置进行辅助对齐。"],
              ["中", "多区域一致性检查", "不要只看一小块是否重合，要检查地图前、中、后三个区域是否同时变好，防止局部对上但整体仍歪。"],
              ["后续", "跨航次位姿图优化", "地图数量很多以后，把各地图之间的重合关系统一优化，而不是每次只把后一张依次接到前面。"],
          ], [1.6, 4.2, 10.2])

add_heading(doc, "十 开会时可以直接这样介绍", 1)
add_body(doc, "我们现在的地图拼接不是直接把两堆点云叠在一起。每份新地图在保存时都有经纬度锚点，并且每个关键帧还保存了对应的 GNSS、INS 的 ENU 参考。加载多份地图后，程序先校正各自内部的慢性漂移，再根据经纬度把它们统一到第一张地图的 ENU 坐标系。考虑到码头里有很多相似的平行结构，自动精配准现在采用保守策略，只允许做几米以内的平移修正，不随便旋转；匹配质量不够就回退到 ENU 结果。重合区域显示时再做三维去重。", bold_lead="我们现在的地图拼接不是直接把两堆点云叠在一起。")
add_body(doc, "为了后面支持更大范围地图，我们又把拼接结果按 50 米网格切成只读小块，并准备了三档精度。界面只加载当前视野和周边区域，远处地图块会自动释放。因此地图范围继续扩大时，不需要每次把整张高分辨率地图全部放进内存。后续优化重点是提高时间同步和静态结构筛选质量，并把配准质量指标显示得更直观。")

add_heading(doc, "十一 演示时建议说明的现象", 1)
for index, text in enumerate([
    "第一次同时选择两张新生成的 .slammap，日志显示地理校正、地图拼接和“已生成只读缓存”。",
    "再次加载相同两张地图，日志显示“已命中只读缓存”，说明不再重新构建全部分块。",
    "切换到俯视图，放大后拖动地图，说明程序正在按当前视野调入相邻 Tile。",
    "滚轮缩放，说明远看使用稀疏 LOD、近看使用更精细 LOD。",
    "强调原始 .slammap 始终只读，分块缓存损坏或删除后可以重新生成。",
], start=1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run(f"{index}. {text}"))

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(footer.add_run("SLAM 地图坐标系统一与地图拼接介绍"), size=8.5, color="777777")

doc.core_properties.title = "SLAM 地图坐标系统一与地图拼接介绍"
doc.core_properties.subject = "当前地图坐标系、拼接算法和只读分块动态显示说明"
doc.core_properties.author = ""
doc.core_properties.keywords = "SLAM, ENU, 地图拼接, 点云, 分块加载"
doc.save(OUTPUT)
print(OUTPUT)
