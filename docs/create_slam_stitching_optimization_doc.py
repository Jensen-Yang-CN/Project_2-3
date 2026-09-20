from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"D:\1\yuanshibanben\PointCloud20260514\SLAM点云地图拼接优化说明.docx")

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
INK = "1F2937"
GREEN = "1F6D4C"
GOLD = "7A5A00"


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_style_font(style, size, color=INK, bold=False):
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (
        ("top", top), ("start", start), ("bottom", bottom), ("end", end)
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_numbering_definition(doc, abstract_id, num_id, ordered=False):
    numbering = doc.part.numbering_part.element
    marker = "%1." if ordered else "•"
    num_fmt = "decimal" if ordered else "bullet"
    xml = f"""
    <w:abstractNum {nsdecls('w')} w:abstractNumId="{abstract_id}">
      <w:multiLevelType w:val="singleLevel"/>
      <w:lvl w:ilvl="0">
        <w:start w:val="1"/>
        <w:numFmt w:val="{num_fmt}"/>
        <w:lvlText w:val="{marker}"/>
        <w:lvlJc w:val="left"/>
        <w:pPr>
          <w:tabs><w:tab w:val="num" w:pos="720"/></w:tabs>
          <w:ind w:left="720" w:hanging="360"/>
          <w:spacing w:after="160" w:line="280" w:lineRule="auto"/>
        </w:pPr>
        <w:rPr>
          <w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:eastAsia="Microsoft YaHei"/>
          <w:sz w:val="22"/>
        </w:rPr>
      </w:lvl>
    </w:abstractNum>
    """
    abstract = parse_xml(xml)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)
    num_xml = f"""
    <w:num {nsdecls('w')} w:numId="{num_id}">
      <w:abstractNumId w:val="{abstract_id}"/>
    </w:num>
    """
    numbering.append(parse_xml(num_xml))


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    apply_numbering(p, 10)
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_run_font(run, bold=True)
        run = p.add_run(text[len(bold_prefix):])
        set_run_font(run)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_numbered(doc, text, num_id=11):
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True, color=NAVY)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        set_run_font(run, bold=True,
                     color=BLUE if level < 3 else DARK_BLUE)
    return p


def set_cell_text(cell, text, bold=False, color=INK, size=10.2):
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_callout(doc, label, text, fill=CALLOUT, accent=NAVY):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    r1 = p.add_run(label)
    set_run_font(r1, size=10.5, bold=True, color=accent)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    set_table_geometry(table, [9360])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.space_before = Pt(2)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, 11, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, 16, BLUE, True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, 13, BLUE, True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, 12, DARK_BLUE, True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    add_numbering_definition(doc, 10, 10, False)
    add_numbering_definition(doc, 11, 11, True)
    add_numbering_definition(doc, 12, 12, True)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("SLAM 地图拼接优化说明  |  PointCloud20260514"),
                 size=9, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(fp.add_run("第 "), size=9, color=MUTED)
    add_field(fp, "PAGE")
    set_run_font(fp.add_run(" 页 / 共 "), size=9, color=MUTED)
    add_field(fp, "NUMPAGES")
    set_run_font(fp.add_run(" 页"), size=9, color=MUTED)


def build_document():
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("技术优化说明")
    set_run_font(r, size=10.5, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("SLAM 点云地图拼接优化说明")
    set_run_font(r, size=24, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("基于逐关键帧地理约束、受限配准与三维显示去重")
    set_run_font(r, size=13, color=MUTED)

    metadata = [
        ("项目", "PointCloud20260514"),
        ("整理日期", "2026 年 8 月 23 日"),
        ("实施状态", "已完成开发、自动化验证与新地图实测"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(f"{label}："), size=10.5, bold=True, color=NAVY)
        set_run_font(p.add_run(value), size=10.5, color=INK)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_callout(
        doc,
        "优化结论：",
        "新程序生成的 v2.1 地图已经能够稳定完成拼接。核心改进不是简单扩大 ICP 搜索范围，而是为每个关键帧保存 GNSS/ENU 参考，在加载阶段先消除地图内部漂移，再使用 3 米以内的保守点云精配准，最后对显示点云进行 0.3 米三维体素去重。",
        fill=LIGHT_BLUE,
        accent=GREEN,
    )

    add_heading(doc, "一、问题背景与根因", 1)
    add_body(
        doc,
        "此前地图拼接虽然能够把多份点云显示在同一界面，但重叠区域常出现红、绿点云双边、轨迹分叉和结构错位。问题并不只是重复点过多，而是旧地图缺少足够的地理约束。",
    )
    add_bullet(doc, "旧版 .slammap v2.0 只保存一个地图地理锚点，只能统一地图原点，不能纠正每份独立 SLAM 在运行过程中产生的累计漂移。")
    add_bullet(doc, "不同 PCAP 独立建图时，GNSS/INS 会话偏差、激光里程计误差和轨迹累计误差会分别进入地图，导致相同结构不能自然重合。")
    add_bullet(doc, "码头、岸线和泊位包含大量相似平行结构，大范围 ICP 容易找到数值上看似合理、实际位置错误的匹配。此前曾出现约 19.7 米的错误候选。")
    add_bullet(doc, "直接删除重复点只能减少显示密度，无法修复坐标或姿态本身的偏差，因此必须先校正位姿，再进行显示去重。")

    heading = add_heading(doc, "二、优化内容总览", 1)
    heading.paragraph_format.page_break_before = True
    rows = [
        ("地图格式", "v2.0 升级为 v2.1", "逐关键帧保存 GNSS/ENU 参考位姿；兼容读取旧地图"),
        ("时间同步", "记录最近原始导航样本误差", "避免插值时间戳导致同步误差被误判为 0"),
        ("内部漂移", "新增鲁棒地理漂移校正器", "先校正每张地图内部轨迹，再进行地图间拼接"),
        ("地图配准", "ICP 限制在 3 米以内", "阻止相似结构产生的大范围错误吸附"),
        ("显示重叠", "0.3 米 XYZ 体素首点保留", "减少重叠区域重复显示，不改动原始地图数据"),
        ("可诊断性", "补充格式、校正、ICP、去重日志", "便于区分旧地图缺少约束、ICP 拒绝和正常校正"),
        ("启动行为", "禁止默认恢复历史地图", "避免旧离线地图与本次解析地图混合，影响判断"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["优化模块", "主要改动", "实际作用"]
    for i, text in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=NAVY)
    set_repeat_table_header(table.rows[0])
    for module, change, effect in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], module, bold=True, color=DARK_BLUE)
        set_cell_text(cells[1], change)
        set_cell_text(cells[2], effect)
    set_table_geometry(table, [1800, 3480, 4080])

    add_heading(doc, "三、核心技术实现", 1)
    add_heading(doc, "3.1 地图格式升级为 v2.1", 2)
    add_body(
        doc,
        "每个 SlamKeyframe 在原 SLAM 位姿之外，新增可选的 SlamGeoPoseReference，保存同一时间戳对应的雷达 ENU 位姿、同步误差和有效标记。新保存的地图统一使用 v2.1；v2.0 仍可读取，但所有逐关键帧地理参考会被标记为无效。高于当前支持版本的文件会明确拒绝，避免按错误布局解析。",
    )
    add_callout(
        doc,
        "数据原则：",
        "SLAM 位姿用于保持局部点云连续性，GNSS/ENU 参考用于离线纠正低频累计漂移。两者同时保存，不用导航瞬时噪声直接覆盖实时 SLAM。",
    )

    heading = add_heading(doc, "3.2 关键帧导航参考同步", 2)
    heading.paragraph_format.page_break_before = True
    add_body(
        doc,
        "雷达帧处理时，系统从 GNSS/INS 里程计中获得匹配位姿，并记录雷达时间戳到最近原始导航样本的真实时间差。创建关键帧时，仅在时间戳一致、位姿有限且同步误差不超过 0.25 秒时写入有效参考。无有效导航信息的关键帧仍可建图，但不会参与地理漂移校正。",
    )

    add_heading(doc, "3.3 逐关键帧鲁棒漂移校正", 2)
    add_body(doc, "每张地图加载后、进入地图间拼接之前，先执行以下处理：")
    add_numbered(doc, "计算每个有效关键帧的二维残差 C_i = T_geo_i × inverse(T_slam_i)。")
    add_numbered(doc, "对平移和偏航残差执行 5 点滑动中值处理，抑制 GNSS 抖动。")
    add_numbered(doc, "相对邻域跳变超过 5 米或 5 度的参考视为离群点；同步误差超过 0.25 秒的参考直接拒绝。")
    add_numbered(doc, "在有效参考之间按时间插值连续修正，参考区间外使用最近端点，不外推漂移趋势。")
    add_numbered(doc, "将平滑修正左乘到关键帧世界位姿，只修改 pose，不修改关键帧局部点云。")
    add_body(
        doc,
        "校正至少需要 3 个可信参考。参考不足时保持原地图位姿不变，并在界面日志中说明原因。这样既能纠正缓慢累计漂移，又不会因为单个 GNSS 跳点扭曲整张地图。",
    )

    add_heading(doc, "3.4 受限地图间配准", 2)
    add_body(
        doc,
        "完成单地图漂移校正后，先利用地理锚点把各地图转换到统一 ENU，再使用小范围 ICP 消除会话级小偏差。默认不改变 ENU 航向，只允许 XY 平移。",
    )
    add_bullet(doc, "最大自动平移修正：3.0 米。")
    add_bullet(doc, "初值搜索步长：1.0 米；距离搜索边界 0.5 米以内的候选拒绝。")
    add_bullet(doc, "必须同时满足双向匹配率、鲁棒 RMSE、绝对改善量和相对改善比例要求。")
    add_bullet(doc, "任何条件不满足均回退到“逐关键帧地理校正 + 锚点 ENU”，不采用未经证明的大幅修正。")

    heading = add_heading(doc, "3.5 三维显示去重", 2)
    heading.paragraph_format.page_break_before = True
    add_body(
        doc,
        "拼接完成后，显示构建器以 0.3 米为边长建立 XYZ 三维体素。同一体素内只显示最先加载地图的第一个点，后续重复点跳过；XY 相同但高度不同的点仍然保留。去重只作用于界面显示缓存，不会从关键帧或保存文件中删除点。",
    )

    add_heading(doc, "四、优化后的完整处理流程", 1)
    flow = doc.add_table(rows=1, cols=1)
    flow.style = "Table Grid"
    flow_cell = flow.cell(0, 0)
    set_cell_shading(flow_cell, CALLOUT)
    flow_p = flow_cell.paragraphs[0]
    flow_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flow_p.paragraph_format.space_before = Pt(6)
    flow_p.paragraph_format.space_after = Pt(6)
    flow_p.paragraph_format.line_spacing = 1.35
    flow_text = (
        "GNSS/INS + 雷达帧时间同步\n"
        "↓\n"
        "SlamKeyframe（SLAM pose + ENU geo_reference）\n"
        "↓\n"
        ".slammap v2.1\n"
        "↓\n"
        "单地图逐关键帧漂移校正\n"
        "↓\n"
        "地理锚点转换到统一 ENU\n"
        "↓\n"
        "3 米以内质量门控 ICP\n"
        "↓\n"
        "0.3 米 XYZ 显示去重\n"
        "↓\n"
        "界面显示与诊断日志"
    )
    set_run_font(flow_p.add_run(flow_text), size=10.8, bold=True, color=NAVY)
    set_table_geometry(flow, [9360])

    add_heading(doc, "五、验证结果", 1)
    verification_rows = [
        ("地图格式测试", "v2.1 保存/读取往返通过；v2.0 兼容读取通过；未来次版本拒绝通过"),
        ("位姿同步测试", "导航参考有效性、时间戳匹配和同步误差规则通过"),
        ("地理锚点测试", "ENU 锚点换算及相关回归测试通过"),
        ("拼接测试", "线性漂移校正、GNSS 跳点、参考不足、3 米边界和 XYZ 去重测试通过"),
        ("旧地图诊断", "3.656 米 ICP 候选超过安全边界，被正确拒绝并回退为锚点对齐"),
        ("程序构建", "PointCloud.exe 完整编译、链接及运行依赖复制成功"),
        ("新地图实测", "使用新程序重新生成地图后，实际拼接效果良好，重叠区域能够正确贴合"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for i, text in enumerate(["验证项目", "结果"]):
        set_cell_shading(table.rows[0].cells[i], LIGHT_GRAY)
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=NAVY)
    set_repeat_table_header(table.rows[0])
    for item, result in verification_rows[:2]:
        cells = table.add_row().cells
        set_cell_text(cells[0], item, bold=True, color=DARK_BLUE)
        set_cell_text(cells[1], result)
    set_table_geometry(table, [2700, 6660])

    page_marker = doc.add_paragraph()
    page_marker.paragraph_format.page_break_before = True
    page_marker.paragraph_format.space_before = Pt(0)
    page_marker.paragraph_format.space_after = Pt(0)
    page_marker.paragraph_format.line_spacing = Pt(1)
    set_run_font(page_marker.add_run(""), size=1, color=WHITE)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for i, text in enumerate(["验证项目", "结果"]):
        set_cell_shading(table.rows[0].cells[i], LIGHT_GRAY)
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=NAVY)
    set_repeat_table_header(table.rows[0])
    for item, result in verification_rows[2:]:
        cells = table.add_row().cells
        set_cell_text(cells[0], item, bold=True, color=DARK_BLUE)
        set_cell_text(cells[1], result)
    set_table_geometry(table, [2700, 6660])

    add_heading(doc, "六、使用方法与注意事项", 1)
    add_numbered(doc, "使用新程序加载原始 PCAP，并让该数据段的 SLAM 建图过程完整运行。", 12)
    add_numbered(doc, "每个 PCAP 分别保存为新的 .slammap 文件；日志应显示格式为 v2.1，并包含有效地理参考。", 12)
    add_numbered(doc, "点击“加载地图”时只选择本次由新程序生成的地图，可使用 Ctrl 多选；选择顺序中的第一张作为基准地图。", 12)
    add_numbered(doc, "检查“地图地理校正”“地图拼接”和“地图显示”日志，确认参考数量、ICP 接受或拒绝原因以及重复点删除数。", 12)
    add_numbered(doc, "需要继续拼接时，优先加载原始的新 v2.1 地图；保存合并结果可用于展示，但不应代替原 PCAP 重新建图。", 12)
    add_callout(
        doc,
        "旧地图注意：",
        "旧 v2.0 文件可以继续加载，但由于没有逐关键帧地理参考，无法从根本恢复其内部漂移。把旧文件简单重新保存成 v2.1 只会改变文件布局，不会补出缺失的地理数据；最高精度必须由原 PCAP 重新生成地图。",
        fill="FFF8E8",
        accent=GOLD,
    )

    add_heading(doc, "七、结论", 1)
    add_body(
        doc,
        "本次优化将地图拼接调整为“逐关键帧地理约束 + 统一 ENU + 小范围质量门控配准 + 三维显示去重”，从位姿、地图间对齐和显示三个层次解决了内部漂移、错误匹配和重复显示。新程序生成的地图已取得良好实测效果，可作为后续地图管理、跨航次更新及栅格/点云双图层开发的基础。",
    )

    doc.core_properties.title = "SLAM 点云地图拼接优化说明"
    doc.core_properties.subject = "PointCloud20260514 地图拼接优化技术报告"
    doc.core_properties.author = "PointCloud 项目组"
    doc.core_properties.keywords = "SLAM, 点云地图, 地图拼接, ENU, GNSS, ICP, 体素去重"
    doc.core_properties.comments = "基于 2026-08-23 完成的地图拼接优化整理"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
